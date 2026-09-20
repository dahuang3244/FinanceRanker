"""XLSX export that reproduces the `Free_Public_Data_Tech_Ranker.xlsx` layout.

Design choice: the workbook writes formulas and lets Excel recalculate. Here the
engine is the source of truth, so every derived cell is written as a *value*.
That keeps the file readable without Excel, avoids stale-formula surprises, and
means the export works headlessly.

Sheet order and headers intentionally match the original template so existing
downstream references keep working.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.config import settings
from app.engine.scoring import METRICS, Z_FIELDS
from app.models import MetricRow

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
NOTE_FONT = Font(italic=True, size=9, color="666666")
NA_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="D9D9D9")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

PCT = "0.0%"
NUM2 = "$0.00"
MULT = '0.00"x"'
MONEY_M = '$#,##0,,"m"'
DATE_FMT = "yyyy-mm-dd"

# Data Cache column layout: (header, attribute, number format)
DATA_CACHE_COLUMNS: list[tuple[str, str, str | None]] = [
    ("Ticker", "ticker", None),
    ("Company", "company", None),
    ("Currency", "currency", None),
    ("Last refresh", "fetched_at", "yyyy-mm-dd hh:mm"),
    ("Status", "status", None),
    ("Price", "price", NUM2),
    ("52W high", "high_52w", NUM2),
    ("52W low", "low_52w", NUM2),
    ("Market cap", "market_cap", MONEY_M),
    ("1Y return", "return_1y", PCT),
    ("Beta", "beta", MULT),
    ("Revenue FY0", "revenue_fy0", MONEY_M),
    ("Revenue growth YoY", "revenue_growth_yoy", PCT),
    ("Revenue CAGR 3Y", "revenue_cagr_3y", PCT),
    ("Revenue CAGR 5Y", "revenue_cagr_5y", PCT),
    ("EPS growth YoY", "eps_growth_yoy", PCT),
    ("Net income growth YoY", "net_income_growth_yoy", PCT),
    ("Gross profit growth YoY", "gross_profit_growth_yoy", PCT),
    ("FCF growth YoY", "fcf_growth_yoy", PCT),
    ("Gross margin", "gross_margin", PCT),
    ("Operating margin", "operating_margin", PCT),
    ("Net margin", "net_margin", PCT),
    ("ROE", "roe", PCT),
    ("ROA", "roa", PCT),
    ("ROIC", "roic", PCT),
    ("Asset turnover", "asset_turnover", MULT),
    ("Operating leverage", "operating_leverage", MULT),
    ("FCF FY0", "fcf_fy0", MONEY_M),
    ("FCF margin", "fcf_margin", PCT),
    ("FCF yield", "fcf_yield", PCT),
    ("OCF / net income", "ocf_to_net_income", MULT),
    ("Debt / assets", "debt_to_assets", PCT),
    ("Capex intensity", "capex_intensity", PCT),
    ("Cash / assets", "cash_to_assets", PCT),
    ("SBC / revenue", "sbc_pct_revenue", PCT),
    ("Net debt FY0", "net_debt_fy0", MONEY_M),
    ("EBITDA FY0", "ebitda_fy0", MONEY_M),
    ("GAAP diluted EPS", "gaap_eps", NUM2),
    ("SBC adj/share", "sbc_adj_share", NUM2),
    ("Restructuring adj/share", "restructuring_adj_share", NUM2),
    ("Amortization adj/share", "amortization_adj_share", NUM2),
    ("Other adj/share", "other_adj_share", NUM2),
    ("Tax adj/share", "tax_adj_share", NUM2),
    ("Model adjusted EPS", "model_adjusted_eps", NUM2),
    ("Reported non-GAAP EPS", "reported_non_gaap_eps", NUM2),
    ("Selected adjusted EPS", "selected_adjusted_eps", NUM2),
    ("EPS FY1 estimate", "eps_fy1_estimate", NUM2),
    ("EPS growth FY1", "eps_growth_fy1", PCT),
    ("Revenue FY1 estimate", "revenue_fy1_estimate", MONEY_M),
    ("Revenue growth FY1", "revenue_growth_fy1", PCT),
    ("GAAP EPS CAGR 5Y", "eps_cagr_5y", PCT),
    ("Forward P/E", "forward_pe", MULT),
    ("PEG", "peg_ratio", MULT),
    ("Price / sales", "price_to_sales", MULT),
    ("EV / sales", "ev_to_sales", MULT),
    ("EV / EBITDA", "ev_to_ebitda", MULT),
    ("Price / book", "price_to_book", MULT),
    ("Price / FCF", "price_to_fcf", MULT),
    ("Price / OCF", "price_to_ocf", MULT),
    ("Net debt / EBITDA", "net_debt_to_ebitda", MULT),
    # --- returns (quoted windows) ---
    ("1M return", "return_1m", PCT),
    ("3M return", "return_3m", PCT),
    ("6M return", "return_6m", PCT),
    ("YTD return", "return_ytd", PCT),
    ("3Y return", "return_3y", PCT),
    # --- returns versus the benchmark ---
    ("SPY 3M return", "benchmark_return_3m", PCT),
    ("SPY 6M return", "benchmark_return_6m", PCT),
    ("SPY 1Y return", "benchmark_return_1y", PCT),
    ("1M excess vs SPY", "excess_return_1m", PCT),
    ("3M excess vs SPY", "excess_return_3m", PCT),
    ("6M excess vs SPY", "excess_return_6m", PCT),
    ("1Y excess vs SPY", "excess_return_1y", PCT),
    ("YTD excess vs SPY", "excess_return_ytd", PCT),
    # --- risk ---
    ("Volatility (ann.)", "volatility", PCT),
    ("Downside deviation", "downside_deviation", PCT),
    ("Sharpe ratio", "sharpe_ratio", MULT),
    ("Sortino ratio", "sortino_ratio", MULT),
    ("Beta (1Y)", "beta_1y", MULT),
    ("Max drawdown 1Y", "max_drawdown_1y", PCT),
    ("52W drawdown", "drawdown_52w", PCT),
    ("Risk obs (days)", "risk_obs_days", "#,##0"),
    ("Risk-free rate used", "risk_free_rate", PCT),
    ("Price basis", "market_price_basis", None),
    # --- consensus ---
    ("Analyst target", "analyst_target", NUM2),
    ("Analyst target median", "analyst_target_median", NUM2),
    ("Analyst target high", "analyst_target_high", NUM2),
    ("Analyst target low", "analyst_target_low", NUM2),
    ("Analyst opinions", "analyst_count", "#,##0"),
    ("Analyst recommendation", "analyst_recommendation", None),
    ("Analyst upside", "analyst_upside", PCT),
    ("Source quality", "source_quality", None),
    ("SEC source", "sec_source", None),
    ("Market source", "market_source", None),
    ("Non-GAAP source", "non_gaap_source", None),
    ("Notes", "notes", None),
    # filing inputs
    ("Revenue FY-1", "revenue_fy_minus_1", MONEY_M),
    ("Revenue FY-5", "revenue_fy_minus_5", MONEY_M),
    ("GAAP EPS FY-1", "gaap_eps_fy_minus_1", NUM2),
    ("GAAP EPS FY-5", "gaap_eps_fy_minus_5", NUM2),
    ("Share count", "share_count", "#,##0"),
    ("Cash FY0", "cash_fy0", MONEY_M),
    ("Debt FY0", "debt_fy0", MONEY_M),
    ("Operating income FY0", "operating_income_fy0", MONEY_M),
    ("Depreciation & amortization FY0", "da_fy0", MONEY_M),
    ("Diluted shares FY0", "diluted_shares_fy0", "#,##0"),
    ("Share compensation FY0", "sbc_fy0", MONEY_M),
    ("Restructuring FY0", "restructuring_fy0", MONEY_M),
    ("Intangible amortization FY0", "amortization_fy0", MONEY_M),
    ("Effective tax rate FY0", "effective_tax_rate", PCT),
    ("Annual period end", "fiscal_end", DATE_FMT),
    ("Equity FY0", "equity_fy0", MONEY_M),
    ("Operating cash flow FY0", "ocf_fy0", MONEY_M),
    ("Capital expenditure FY0", "capex_fy0", MONEY_M),
    ("Capex basis", "capex_basis", None),
    ("Debt basis", "debt_basis", None),
    ("Margin basis", "margin_basis", None),
    ("Share-count basis", "share_basis", None),
    ("Total assets FY0", "total_assets_fy0", MONEY_M),
    ("Gross profit FY0", "gross_profit_fy0", MONEY_M),
    ("Cost of revenue FY0", "cost_of_revenue_fy0", MONEY_M),
    ("Forward 12m EPS", "forward_12m_eps", NUM2),
    ("Quote checked", "fetched_at", "yyyy-mm-dd hh:mm"),
    ("Quote URL", "quote_url", None),
]

# Peer Data metric rows: (label, unit, attribute)
PEER_METRICS: list[tuple[str, str, str]] = [
    ("Company name", "text", "company"),
    ("Currency", "text", "currency"),
    ("Last refresh", "date", "fetched_at"),
    ("Status", "text", "status"),
    ("Price", "USD", "price"),
    ("52W high", "USD", "high_52w"),
    ("52W low", "USD", "low_52w"),
    ("Market cap", "USD", "market_cap"),
    ("1Y return", "%", "return_1y"),
    ("3M return", "%", "return_3m"),
    ("6M return", "%", "return_6m"),
    ("1M return", "%", "return_1m"),
    ("YTD return", "%", "return_ytd"),
    ("3M excess vs SPY", "%", "excess_return_3m"),
    ("6M excess vs SPY", "%", "excess_return_6m"),
    ("1Y excess vs SPY", "%", "excess_return_1y"),
    ("SPY 6M return", "%", "benchmark_return_6m"),
    ("Volatility (ann.)", "%", "volatility"),
    ("Sharpe ratio", "x", "sharpe_ratio"),
    ("Sortino ratio", "x", "sortino_ratio"),
    ("Beta", "x", "beta"),
    ("Beta (1Y)", "x", "beta_1y"),
    ("Max drawdown 1Y", "%", "max_drawdown_1y"),
    ("Revenue CAGR 3Y", "%", "revenue_cagr_3y"),
    ("EPS growth YoY", "%", "eps_growth_yoy"),
    ("Net income growth YoY", "%", "net_income_growth_yoy"),
    ("Gross profit growth YoY", "%", "gross_profit_growth_yoy"),
    ("FCF growth YoY", "%", "fcf_growth_yoy"),
    ("ROA", "%", "roa"),
    ("Asset turnover", "x", "asset_turnover"),
    ("Operating leverage", "x", "operating_leverage"),
    ("Capex intensity", "%", "capex_intensity"),
    ("Cash / assets", "%", "cash_to_assets"),
    ("SBC / revenue", "%", "sbc_pct_revenue"),
    ("Price / OCF", "x", "price_to_ocf"),
    ("EV / sales", "x", "ev_to_sales"),
    ("PEG", "x", "peg_ratio"),
    ("Net debt / EBITDA", "x", "net_debt_to_ebitda"),
    ("Revenue FY0", "USD", "revenue_fy0"),
    ("Revenue growth YoY", "%", "revenue_growth_yoy"),
    ("Revenue CAGR 5Y", "%", "revenue_cagr_5y"),
    ("Gross margin", "%", "gross_margin"),
    ("Operating margin", "%", "operating_margin"),
    ("Net margin", "%", "net_margin"),
    ("ROE", "%", "roe"),
    ("ROIC", "%", "roic"),
    ("FCF FY0", "USD", "fcf_fy0"),
    ("FCF margin", "%", "fcf_margin"),
    ("FCF yield", "%", "fcf_yield"),
    ("OCF / net income", "x", "ocf_to_net_income"),
    ("Debt / assets", "%", "debt_to_assets"),
    ("GAAP diluted EPS", "USD/share", "gaap_eps"),
    ("Model adjusted EPS", "USD/share", "model_adjusted_eps"),
    ("Reported non-GAAP EPS", "USD/share", "reported_non_gaap_eps"),
    ("Selected adjusted EPS", "USD/share", "selected_adjusted_eps"),
    ("EPS FY1 estimate", "USD/share", "eps_fy1_estimate"),
    ("EPS growth FY1", "%", "eps_growth_fy1"),
    ("Revenue FY1 estimate", "USD", "revenue_fy1_estimate"),
    ("Revenue growth FY1", "%", "revenue_growth_fy1"),
    ("GAAP EPS CAGR 5Y", "%", "eps_cagr_5y"),
    ("Forward P/E", "x", "forward_pe"),
    ("Price / sales", "x", "price_to_sales"),
    ("EV / EBITDA", "x", "ev_to_ebitda"),
    ("Price / book", "x", "price_to_book"),
    ("Price / FCF", "x", "price_to_fcf"),
    ("52W drawdown", "%", "drawdown_52w"),
    ("Analyst target", "USD", "analyst_target"),
    ("Analyst upside", "%", "analyst_upside"),
    ("Analyst opinions", "count", "analyst_count"),
    ("Analyst recommendation", "text", "analyst_recommendation"),
    ("Price basis", "text", "market_price_basis"),
    ("Source quality", "text", "source_quality"),
]

SCORING_ROWS: list[tuple[str, str, str]] = [
    ("Revenue growth YoY", "growth", "z_growth_yoy"),
    ("Revenue CAGR 5Y", "growth", "z_revenue_cagr"),
    ("FY1 EPS growth", "growth", "z_eps_growth_fy1"),
    ("GAAP EPS CAGR 5Y", "growth", "z_eps_cagr"),
    ("Gross margin", "profitability", "z_gross_margin"),
    ("Operating margin", "profitability", "z_operating_margin"),
    ("Net margin", "profitability", "z_net_margin"),
    ("ROE", "profitability", "z_roe"),
    ("ROIC", "profitability", "z_roic"),
    ("FCF margin", "cash", "z_fcf_margin"),
    ("FCF yield", "cash", "z_fcf_yield"),
    ("Cash conversion", "cash", "z_cash_conversion"),
    ("Debt / assets", "cash", "z_debt_assets"),
    ("Forward P/E", "valuation", "z_forward_pe"),
    ("Price / sales", "valuation", "z_price_sales"),
    ("EV / EBITDA", "valuation", "z_ev_to_ebitda"),
    ("Price / book", "valuation", "z_price_book"),
    ("Price / FCF", "valuation", "z_price_fcf"),
    ("3M return", "market", "z_return_3m"),
    ("6M return", "market", "z_return_6m"),
    ("1Y return", "market", "z_return_1y"),
    ("3M excess vs SPY", "market", "z_excess_return_3m"),
    ("6M excess vs SPY", "market", "z_excess_return_6m"),
    ("1Y excess vs SPY", "market", "z_excess_return_1y"),
    ("Sharpe ratio", "market", "z_sharpe"),
    ("Sortino ratio", "market", "z_sortino"),
    ("Volatility", "market", "z_volatility"),
    ("Beta (1Y)", "market", "z_beta_1y"),
    ("Max drawdown 1Y", "market", "z_max_drawdown"),
    ("52W drawdown", "market", "z_drawdown"),
    ("ROA", "profitability", "z_roa"),
    ("Capex intensity", "cash", "z_capex_intensity"),
]


def _number_format_for(attr: str) -> str | None:
    for _, key, fmt in DATA_CACHE_COLUMNS:
        if key == attr:
            return fmt
    return None


def _style_header(ws, row: int, count: int) -> None:
    for col in range(1, count + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = CELL_BORDER


def _autosize(ws, max_width: int = 42) -> None:
    for column in ws.columns:
        length = 0
        letter = get_column_letter(column[0].column)
        for cell in column:
            if cell.value is not None:
                length = max(length, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(10, length + 2), max_width)


def _sheet_ranking(wb: Workbook, rows: list[MetricRow]) -> None:
    ws = wb.create_sheet("Ranking")
    ws["A2"] = "Technology peer ranking — free public data"
    ws["A2"].font = TITLE_FONT
    ws["A4"] = "Last refresh"
    ws["B4"] = datetime.now()
    ws["B4"].number_format = "yyyy-mm-dd hh:mm"
    ws["A5"] = "Source mode"
    ws["B5"] = "SEC XBRL / akshare fundamentals + public price history"
    ws["C4"] = "Ranked peers"
    ws["D4"] = sum(1 for r in rows if r.rank is not None)
    ws["C5"] = "Scoring range"
    ws["D5"] = "1 to 10"

    headers = [
        "Rank", "Ticker", "Company", "Growth", "Profitability", "Cash quality",
        "Valuation", "Market / risk", "Overall", "Profile", "Data coverage", "Status",
    ]
    ws.append([])
    ws.append(headers)
    header_row = ws.max_row
    _style_header(ws, header_row, len(headers))

    for row in rows:
        ws.append([
            row.rank, row.ticker, row.company, row.score_growth, row.score_profitability,
            row.score_cash, row.score_valuation, row.score_market, row.score_overall,
            row.profile, f"{row.data_coverage} / {len(METRICS)}",
            row.status,
        ])
        current = ws.max_row
        for col in range(4, 10):
            ws.cell(current, col).number_format = "0.00"
    _autosize(ws)
    ws.freeze_panes = ws.cell(header_row + 1, 3)


def _sheet_peer_data(wb: Workbook, rows: list[MetricRow]) -> None:
    ws = wb.create_sheet("Peer Data")
    ws["A2"] = "Automated peer data — free public sources"
    ws["A2"].font = TITLE_FONT
    ws["A4"] = "Last refresh"
    ws["B4"] = datetime.now()
    ws["B4"].number_format = "yyyy-mm-dd hh:mm"
    ws["A5"] = "Ticker inputs"
    for idx, row in enumerate(rows):
        ws.cell(5, 4 + idx, row.ticker)

    start_row = 7
    for offset, (label, unit, _) in enumerate(PEER_METRICS):
        r = start_row + offset
        ws.cell(r, 1, label)
        ws.cell(r, 3, unit)
    for idx, row in enumerate(rows):
        col = 4 + idx
        for offset, (_, _, attr) in enumerate(PEER_METRICS):
            r = start_row + offset
            value = getattr(row, attr, None)
            cell = ws.cell(r, col, value if value is not None else "Public N/A")
            fmt = _number_format_for(attr)
            if fmt and value is not None:
                cell.number_format = fmt
            elif value == "Public N/A":
                cell.fill = NA_FILL
    _autosize(ws, 30)


def _sheet_data_cache(wb: Workbook, rows: list[MetricRow]) -> None:
    ws = wb.create_sheet("Data Cache")
    ws["A2"] = "Public-data cache"
    ws["A2"].font = TITLE_FONT
    for idx, (header, attr, fmt) in enumerate(DATA_CACHE_COLUMNS, start=1):
        ws.cell(4, idx, header)
        ws.cell(5, idx, attr)
        cell = ws.cell(4, idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, row in enumerate(rows, start=6):
        for idx, (_, attr, fmt) in enumerate(DATA_CACHE_COLUMNS, start=1):
            value = getattr(row, attr, None)
            cell = ws.cell(r, idx, value if value is not None else "")
            if fmt and value is not None:
                cell.number_format = fmt
    _autosize(ws, 26)
    ws.freeze_panes = "C6"


def _sheet_scoring(wb: Workbook, rows: list[MetricRow]) -> None:
    ws = wb.create_sheet("Scoring")
    ws["A2"] = "Cross-sectional percentile scoring (1–10)"
    ws["A2"].font = TITLE_FONT
    ws["A4"] = "Weights"
    ws["B4"] = settings.w_growth
    ws["C4"] = settings.w_profitability
    ws["D4"] = settings.w_cash
    ws["E4"] = settings.w_valuation
    ws["F4"] = settings.w_market
    for col in range(2, 7):
        ws.cell(4, col).number_format = PCT

    ws.append([])
    ws.append(["Ticker", "Metric", "Component"])
    header_row = ws.max_row
    _style_header(ws, header_row, 3 + len(rows))
    for idx in range(len(rows)):
        ws.cell(header_row, 4 + idx, rows[idx].ticker)

    r = header_row
    for label, component, attr in SCORING_ROWS:
        r += 1
        ws.cell(r, 1, f"z:{label}")
        ws.cell(r, 2, label)
        ws.cell(r, 3, component)
        for idx, row in enumerate(rows):
            cell = ws.cell(r, 4 + idx, getattr(row, attr, None))
            cell.number_format = "0.00"

    for label, attr in (
        ("Score growth", "score_growth"),
        ("Score profitability", "score_profitability"),
        ("Score cash", "score_cash"),
        ("Score valuation", "score_valuation"),
        ("Score market", "score_market"),
        ("Score overall", "score_overall"),
    ):
        r += 1
        ws.cell(r, 1, label)
        for idx, row in enumerate(rows):
            cell = ws.cell(r, 4 + idx, getattr(row, attr, None))
            cell.number_format = "0.00"
    _autosize(ws, 24)


def _sheet_non_gaap_bridge(wb: Workbook, rows: list[MetricRow]) -> None:
    ws = wb.create_sheet("Non-GAAP Bridge")
    ws["A2"] = "Non-GAAP EPS bridge"
    ws["A2"].font = TITLE_FONT
    headers = [
        "Ticker", "GAAP diluted EPS", "SBC adj/share", "Restructuring adj/share",
        "Amortization adj/share", "Other adj/share", "Tax adj/share",
        "Model adjusted EPS", "Reported non-GAAP EPS", "Selected adjusted EPS",
        "Adjustment %", "Source status", "Non-GAAP source", "Review note",
    ]
    ws.append([])
    ws.append(headers)
    header_row = ws.max_row
    _style_header(ws, header_row, len(headers))
    for row in rows:
        base = row.gaap_eps
        selected = row.selected_adjusted_eps
        adjustment = (selected / base - 1.0) if (base and selected) else None
        ws.append([
            row.ticker, row.gaap_eps, row.sbc_adj_share, row.restructuring_adj_share,
            row.amortization_adj_share, row.other_adj_share, row.tax_adj_share,
            row.model_adjusted_eps, row.reported_non_gaap_eps, selected,
            adjustment,
            "Company reported" if row.reported_non_gaap_eps else (
                "Model estimate" if row.model_adjusted_eps else "No comparable EPS"
            ),
            row.non_gaap_source or "Public N/A",
            row.notes,
        ])
        current = ws.max_row
        for col in range(2, 11):
            ws.cell(current, col).number_format = NUM2
        ws.cell(current, 11).number_format = PCT
    _autosize(ws)
    ws.freeze_panes = ws.cell(header_row + 1, 2)


def _sheet_stock_detail(wb: Workbook, rows: list[MetricRow]) -> None:
    ws = wb.create_sheet("Stock Detail")
    ws["A2"] = "Single-stock peer review"
    ws["A2"].font = TITLE_FONT
    ws["A4"] = "Selected ticker"
    ws["B4"] = rows[0].ticker if rows else ""
    ws["D4"] = "Overall score"
    ws["E4"] = rows[0].score_overall if rows else ""
    ws["F4"] = "Peer rank"
    ws["G4"] = rows[0].rank if rows else ""

    ws["A6"] = "Component"
    for idx, label in enumerate(
        ["Growth", "Profitability", "Cash quality", "Valuation", "Market / risk", "Overall"], start=2
    ):
        ws.cell(6, idx, label)
    ws["A7"] = "Score (1–10)"
    if rows:
        top = rows[0]
        for idx, value in enumerate(
            [top.score_growth, top.score_profitability, top.score_cash,
             top.score_valuation, top.score_market, top.score_overall], start=2
        ):
            ws.cell(7, idx, value).number_format = "0.00"

    ws["A10"] = "Metric"
    ws["B10"] = "Stock"
    ws["C10"] = "Peer median"
    ws["D10"] = "Score"
    metrics = [
        ("Revenue growth YoY", "revenue_growth_yoy", PCT),
        ("Revenue CAGR 5Y", "revenue_cagr_5y", PCT),
        ("Gross margin", "gross_margin", PCT),
        ("Operating margin", "operating_margin", PCT),
        ("Net margin", "net_margin", PCT),
        ("ROE", "roe", PCT),
        ("ROIC", "roic", PCT),
        ("FCF margin", "fcf_margin", PCT),
        ("FCF yield", "fcf_yield", PCT),
        ("Forward P/E", "forward_pe", MULT),
        ("Price / sales", "price_to_sales", MULT),
        ("EV / EBITDA", "ev_to_ebitda", MULT),
        ("Price / book", "price_to_book", MULT),
        ("Price / FCF", "price_to_fcf", MULT),
        ("1Y return", "return_1y", PCT),
        ("3M return", "return_3m", PCT),
        ("6M excess vs SPY", "excess_return_6m", PCT),
        ("Volatility (ann.)", "volatility", PCT),
        ("Sharpe ratio", "sharpe_ratio", MULT),
        ("Max drawdown 1Y", "max_drawdown_1y", PCT),
        ("Debt / assets", "debt_to_assets", PCT),
        ("52W drawdown", "drawdown_52w", PCT),
        ("Beta", "beta", MULT),
        ("Beta (1Y)", "beta_1y", MULT),
    ]
    if rows:
        top = rows[0]
        r = 10
        for label, attr, fmt in metrics:
            r += 1
            ws.cell(r, 1, label)
            values = [getattr(x, attr) for x in rows if getattr(x, attr) is not None]
            stock_value = getattr(top, attr, None)
            ws.cell(r, 2, stock_value).number_format = fmt
            if values:
                ordered = sorted(values)
                mid = len(ordered) // 2
                median = (
                    ordered[mid]
                    if len(ordered) % 2
                    else (ordered[mid - 1] + ordered[mid]) / 2
                )
                ws.cell(r, 3, median).number_format = fmt
            z_field = Z_FIELDS.get(attr)
            z_value = getattr(top, z_field, None) if z_field else None
            if z_value is not None:
                ws.cell(r, 4, z_value).number_format = "0.00"
    _autosize(ws)


def _sheet_guidance(wb: Workbook, rows: list[MetricRow]) -> None:
    ws = wb.create_sheet("Guidance")
    ws["A1"] = "Company guidance: full-year EPS only"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (
        "Optional sheet. Fill in a sourced, published full-year EPS range to override the "
        "provider estimate used for FY1 EPS growth."
    )
    ws["A2"].font = NOTE_FONT
    headers = [
        "Ticker", "FY end year", "Release date", "EPS basis", "Low EPS", "High EPS",
        "Source URL", "Midpoint EPS", "Coverage note",
    ]
    ws.append([])
    ws.append(headers)
    header_row = ws.max_row
    _style_header(ws, header_row, len(headers))
    for row in rows:
        ws.append([
            row.ticker,
            row.fiscal_end.year if row.fiscal_end else None,
            None, None, None, None, None, None,
            "Optional: full-year guidance not yet captured",
        ])
    _autosize(ws)


def _sheet_methodology(wb: Workbook, rows: list[MetricRow]) -> None:
    ws = wb.create_sheet("Methodology")
    ws["A2"] = "Methodology & data provenance"
    ws["A2"].font = TITLE_FONT
    lines = [
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Peers", str(len(rows))),
        ("Eligible for ranking", str(sum(1 for r in rows if r.rank_eligible))),
        ("", ""),
        ("Fundamentals", "SEC XBRL companyfacts (annual 10-K / 20-F / 40-F), akshare fallback"),
        ("Prices / quotes", "Sina US daily bars, Tencent realtime quote, Eastmoney / Yahoo fallback"),
        ("Fiscal anchoring", "Newest annual period across all candidate XBRL tags"),
        ("Non-GAAP bridge", "After-tax per-share add-backs: SBC, restructuring, intangible amortization"),
        ("ROIC", "Operating income x (1 - effective tax) / (equity + debt - cash)"),
        ("EV / EBITDA", "Market cap + debt - cash, over operating income + D&A"),
        ("Beta", "Covariance / variance of aligned daily returns vs SPY (min 30 observations)"),
        ("Scoring", "1-10 percentile rank per metric, averaged per component, weighted 25/25/20/20/10"),
        ("", ""),
        ("Caveats", ""),
        ("Price", "Derived from daily closes (adjusted where available); intraday extremes may differ"),
        ("FY1 estimates", "Sourced from provider consensus when available; otherwise not scored"),
        ("Non-GAAP EPS", "Not auto-sourced; the bridge is a model estimate, not company guidance"),
        ("Use", "Research screen only — verify against filings before any investment decision"),
    ]
    r = 4
    for label, value in lines:
        ws.cell(r, 1, label).font = Font(bold=True)
        ws.cell(r, 2, value)
        r += 1
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 96


def build_workbook(rows: list[MetricRow]) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)  # drop the default sheet
    _sheet_ranking(wb, rows)
    _sheet_stock_detail(wb, rows)
    _sheet_peer_data(wb, rows)
    _sheet_scoring(wb, rows)
    _sheet_non_gaap_bridge(wb, rows)
    _sheet_data_cache(wb, rows)
    _sheet_methodology(wb, rows)
    _sheet_guidance(wb, rows)
    return wb


def write_xlsx(rows: list[MetricRow], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(rows).save(path)
    return path
